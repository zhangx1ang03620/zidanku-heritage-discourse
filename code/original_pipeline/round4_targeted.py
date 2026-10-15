#!/usr/bin/env python3
"""
round4_targeted.py —— 定向验证（第 4 轮，193 条）的回收检查、分歧表、合成与分析

  python3 round4_targeted.py --check      格式检查 + 两人 κ + 导出 裁决表_第4轮.xlsx（不读层标记，不读模型标签）
  python3 round4_targeted.py --analyze    裁决填完后：合成人工答案，按日志 7.12 预先写定的估计量和判定规则出结果

复用 round3_merge.py 的检查与 κ 函数。
"""
import argparse
import os
import sys

import numpy as np
import pandas as pd

import round3_merge as R

BASE = R.BASE
R4 = os.path.join(BASE, "人工编码任务包", "第4轮_定向验证")
R.R3, R.BACK = R4, os.path.join(R4, "回收")
R.FILES = {"r4": "4_补充193条"}
ADJ = os.path.join(R4, "裁决表_第4轮.xlsx")
ADJ_KEY = os.path.join(BASE, "db", "targeted_validation_adj_key.csv")
GOLD = os.path.join(BASE, "db", "targeted_validation_human.csv")
TABLES = os.path.join(os.path.dirname(BASE), "8.分析与图表", "tables")
C, KEY = R.COLS, R.KEY
S = "质疑异见"
B = 2000


def export_adj(m):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    d = m.fillna("")
    dis = {c: d[f"{c}_A"] != d[f"{c}_B"] for c in C}
    d = d[dis[C[0]] | dis[C[1]] | dis[C[2]]].copy()
    d["分歧项"] = ["、".join(n for c, n in zip(C, ["相关", "语气", "功能"]) if dis[c][i]) for i in d.index]
    # 裁决表不显示平台和 comment_id：负责人知道本轮在检验哪个平台的哪一类，隐去平台以免影响裁决
    d = d.rename(columns={"帖子主题_A": "帖子主题", "评论内容_A": "评论内容"}).sample(frac=1, random_state=20261011)
    d = d.sort_values("分歧项", kind="stable")
    d["编号"] = [f"R{i + 1:03d}" for i in range(len(d))]
    d[["编号"] + KEY].to_csv(ADJ_KEY, index=False, encoding="utf-8-sig")
    cols = ["编号", "帖子主题", "评论内容", "分歧项"] + [f"{c}_{w}" for c in C for w in "AB"] + ["备注_A", "备注_B"] + [f"裁决_{c}" for c in C]
    wb = Workbook()
    ws = wb.active
    ws.title = "裁决"
    ws.append(cols)
    for r in d[cols[:-3]].itertuples(index=False):
        ws.append(list(r) + ["", "", ""])
    n = len(d) + 1
    for i in range(3):
        for r in range(2, n + 1):
            a, b = ws.cell(r, 5 + 2 * i), ws.cell(r, 6 + 2 * i)
            if a.value != b.value:
                a.font = b.font = Font(color="C00000", bold=True)
            ws.cell(r, 13 + i).fill = PatternFill("solid", fgColor="FFF2CC")
    for col, vals in zip("MNO", [["是", "否"], R.TONES, R.FUNCTIONS]):
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    for col, w in zip("ABCDEFGHIJKLMNO", [8, 40, 50, 12, 8, 8, 10, 10, 18, 18, 22, 22, 10, 12, 20]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.number_format = "@"
    ws.freeze_panes = "D2"
    wb.save(ADJ)
    print(f"\n两人一致 {len(m) - len(d)} 条，分歧 {len(d)} 条（{len(d) / len(m):.1%}）：{d.分歧项.value_counts().to_dict()}\n→ {ADJ}")


def gold():
    m = R.pair("r4")
    adj = pd.read_excel(ADJ, sheet_name=0, dtype=str)
    adj = adj.apply(lambda c: c.str.strip()).replace({"": np.nan})
    adj = adj.merge(pd.read_csv(ADJ_KEY, dtype=str, encoding="utf-8-sig"), on="编号").set_index(KEY)
    mi = m.set_index(KEY)
    g = m[KEY].copy()
    for c in C:
        v = mi[f"{c}_A"].where(mi[f"{c}_A"].fillna("") == mi[f"{c}_B"].fillna(""))
        v.update(adj[f"裁决_{c}"].dropna())
        g[c] = v.to_numpy()
    g["来源"] = np.where(mi.index.isin(adj.index), "裁决", "两人一致")
    yes = g.human_on_topic == "是"
    missing = int(g.human_on_topic.isna().sum()) + int((yes & (g.human_tone.isna() | g.human_function.isna())).sum())
    if missing:
        sys.exit(f"裁决表未填完：缺 {missing} 格。")
    g.loc[~yes, ["human_tone", "human_function"]] = np.nan
    g.to_csv(GOLD, index=False, encoding="utf-8-sig")
    return g


def estimate(d, strata):
    """各平台校正占比 = π_S·PPV + (1−π_S)·FOR；d 含 platform, stratum, hs（人工是否判相关且质疑）"""
    r = {}
    tot = {}
    for p, g in strata.groupby("platform"):
        pop = g.set_index("stratum").population
        pi = pop["S"] / pop.sum()
        ppv = d[(d.platform == p) & (d.stratum == "S")].hs.mean()
        fo = d[(d.platform == p) & (d.stratum == "N")].hs.mean()
        r[f"{p}|PPV"], r[f"{p}|FOR"], r[f"{p}|share_model"] = ppv, fo, pi
        r[f"{p}|share_corrected"] = pi * ppv + (1 - pi) * fo
        tot[p] = (pop.sum(), r[f"{p}|share_corrected"])
    others = [p for p in tot if p != "zhihu"]
    rest = sum(tot[p][0] * tot[p][1] for p in others) / sum(tot[p][0] for p in others)
    r["others_pooled|share_corrected"] = rest
    r["zhihu_minus_others"] = tot["zhihu"][1] - rest
    r["zhihu_is_highest"] = float(tot["zhihu"][1] == max(v[1] for v in tot.values()))
    return r


def analyze():
    g = gold()
    fr = pd.read_csv(os.path.join(BASE, "db", "targeted_validation_frame.csv"), dtype=str)
    strata = pd.read_csv(os.path.join(BASE, "db", "targeted_validation_strata.csv"))
    d = g.merge(fr, on=KEY)
    d["hs"] = ((d.human_on_topic == "是") & (d.human_tone == S)).astype(float)
    point = estimate(d, strata)
    rng = np.random.default_rng(20261010)
    groups = [x for _, x in d.groupby(["platform", "stratum"])]
    bt = pd.DataFrame([estimate(pd.concat([x.iloc[rng.integers(0, len(x), len(x))] for x in groups]), strata) for _ in range(B)])
    t = pd.DataFrame({"metric": list(point), "estimate": list(point.values()),
                      "ci_low": [np.nanpercentile(bt[k], 2.5) for k in point], "ci_high": [np.nanpercentile(bt[k], 97.5) for k in point]})
    t.loc[t.metric.str.endswith("share_model") | (t.metric == "zhihu_is_highest"), ["ci_low", "ci_high"]] = np.nan
    t.to_csv(os.path.join(TABLES, "T16_targeted_validation_skepticism.csv"), index=False, encoding="utf-8-sig")
    print(t.round(3).to_string(index=False))
    print("知乎在 bootstrap 中为最高的比例：", round(bt.zhihu_is_highest.mean(), 3))
    diff = t.set_index("metric").loc["zhihu_minus_others"]
    keep = point["zhihu_is_highest"] == 1 and diff.ci_low > 0
    print("\n判定（日志 7.12）：" + ("知乎校正占比最高且差值区间不含 0 → 平台 × 质疑异见保留为主要发现。" if keep
                               else "未满足预先写定的条件 → 平台 × 质疑异见降为探索性。"))
    print(pd.crosstab([d.platform, d.stratum], d.human_tone.fillna("（人工判不相关）")).to_string())

    # 对裁决的敏感性：只用 A、只用 B 的答案各算一遍（日志 7.12 回收后追加）
    m = R.pair("r4").merge(fr, on=KEY)
    rows = []
    for w in "AB":
        x = m.assign(hs=((m[f"human_on_topic_{w}"] == "是") & (m[f"human_tone_{w}"] == S)).astype(float))
        pt = estimate(x, strata)
        gs = [y for _, y in x.groupby(["platform", "stratum"])]
        b2 = pd.DataFrame([estimate(pd.concat([y.iloc[rng.integers(0, len(y), len(y))] for y in gs]), strata) for _ in range(B)])
        lo, hi = np.percentile(b2.zhihu_minus_others, [2.5, 97.5])
        rows.append({"answers": f"only {w}", **{k: v for k, v in pt.items() if k.endswith("share_corrected") or k.endswith("PPV")},
                     "zhihu_minus_others": pt["zhihu_minus_others"], "diff_ci_low": lo, "diff_ci_high": hi,
                     "zhihu_is_highest": pt["zhihu_is_highest"], "keep_primary": bool(pt["zhihu_is_highest"] == 1 and lo > 0)})
    sens = pd.DataFrame(rows)
    sens.to_csv(os.path.join(TABLES, "T16b_targeted_validation_sensitivity.csv"), index=False, encoding="utf-8-sig")
    print("\n敏感性：只用一名编码者的答案\n", sens.round(3).T.to_string())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--analyze", action="store_true")
    a = ap.parse_args()
    if a.check:
        m = R.pair("r4")
        R.report_kappa(m)
        export_adj(m)
    elif a.analyze:
        analyze()
    else:
        ap.print_help()
